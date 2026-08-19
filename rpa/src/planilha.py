"""Leitura da planilha do desafio com OpenPyXL.

Responsabilidades: localizar o arquivo, validar a estrutura, normalizar os
cabeçalhos e transformar as linhas em registros. Este módulo não conhece
Selenium nem o formulário web.
"""

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .erros import ErroPlanilha

logger = logging.getLogger(__name__)

# Cabeçalho normalizado -> campo do registro.
# A planilha oficial traz "Last Name " com espaço à direita, por isso o
# cabeçalho é normalizado antes do mapeamento.
COLUNAS = {
    "first name": "first_name",
    "last name": "last_name",
    "company name": "company_name",
    "role in company": "role_in_company",
    "address": "address",
    "email": "email",
    "phone number": "phone_number",
}


@dataclass(frozen=True)
class Registro:
    """Uma linha da planilha, pronta para ser digitada no formulário."""

    first_name: str
    last_name: str
    company_name: str
    role_in_company: str
    address: str
    email: str
    phone_number: str


def normalizar_cabecalho(valor: object) -> str:
    """Normaliza um cabeçalho: colapsa espaços e remove diferença de caixa."""
    if valor is None:
        return ""
    return " ".join(str(valor).split()).lower()


def normalizar_valor(valor: object) -> str:
    """Converte o valor da célula em texto preservando o dado original.

    Telefones chegam como int e não podem virar "4.0716543298e+10"; floats
    inteiros (40716543298.0) perdem o ".0"; demais tipos viram str direto.
    """
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else repr(valor)
    return str(valor).strip()


def _mapear_colunas(cabecalho: tuple) -> dict[str, int]:
    """Relaciona cada campo esperado ao índice da coluna correspondente."""
    posicoes: dict[str, int] = {}
    for indice, celula in enumerate(cabecalho):
        campo = COLUNAS.get(normalizar_cabecalho(celula))
        if campo and campo not in posicoes:
            posicoes[campo] = indice

    ausentes = sorted(set(COLUNAS.values()) - set(posicoes))
    if ausentes:
        encontrados = [normalizar_cabecalho(c) for c in cabecalho if c is not None]
        raise ErroPlanilha(
            f"colunas obrigatórias ausentes na planilha: {ausentes}. "
            f"Cabeçalhos encontrados: {encontrados}"
        )
    return posicoes


def ler_registros(caminho: Path) -> list[Registro]:
    """Lê a planilha e devolve os registros na ordem em que aparecem."""
    if not caminho.is_file():
        raise ErroPlanilha(f"planilha não encontrada em {caminho}")

    planilha = load_workbook(caminho, data_only=True, read_only=True)
    try:
        aba = planilha.active
        linhas = aba.iter_rows(values_only=True)
        cabecalho = next(linhas, None)
        if cabecalho is None:
            raise ErroPlanilha(f"planilha {caminho.name} está vazia")

        posicoes = _mapear_colunas(cabecalho)

        registros: list[Registro] = []
        for numero, linha in enumerate(linhas, start=2):
            valores = {
                campo: normalizar_valor(linha[indice]) if indice < len(linha) else ""
                for campo, indice in posicoes.items()
            }
            if not any(valores.values()):
                continue  # a planilha oficial traz centenas de linhas em branco
            vazios = sorted(campo for campo, valor in valores.items() if not valor)
            if vazios:
                raise ErroPlanilha(
                    f"linha {numero} da planilha tem campos obrigatórios vazios: {vazios}"
                )
            registros.append(Registro(**valores))
    finally:
        planilha.close()

    if not registros:
        raise ErroPlanilha(f"planilha {caminho.name} não contém registros")

    logger.info("Planilha lida: %d registros em %s", len(registros), caminho.name)
    return registros
